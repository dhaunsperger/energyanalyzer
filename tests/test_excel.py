"""Tests for report/excel.py (ARCHITECTURE.md §9 Task 6 — Excel builder only)."""

from __future__ import annotations

import datetime as dt
import io

import openpyxl
import pandas as pd
import pytest

from energyanalyzer.core.models import Buyback, BuybackKind, EnergyRate, Plan, TduTariff
from energyanalyzer.engine.cost import rank
from energyanalyzer.report.excel import build_workbook


def make_intervals(start: str, days: int, import_kwh: float, export_kwh: float) -> pd.DataFrame:
    """Synthetic canonical interval frame: constant import/export per 15-min
    interval, `days` days starting at local midnight on `start`."""
    idx_local = pd.date_range(start=start, periods=days * 96, freq="15min", tz="America/Chicago")
    idx_utc = idx_local.tz_convert("UTC")
    return pd.DataFrame(
        {"import_kwh": import_kwh, "export_kwh": export_kwh}, index=idx_utc
    ).sort_index()


@pytest.fixture
def intervals() -> pd.DataFrame:
    # Two full calendar months (Jan, Feb 2024) of constant usage, avoiding
    # DST transitions so localization is unambiguous.
    jan = make_intervals("2024-01-01", days=31, import_kwh=0.5, export_kwh=0.2)
    feb = make_intervals("2024-02-01", days=29, import_kwh=0.4, export_kwh=0.3)
    return pd.concat([jan, feb]).sort_index()


@pytest.fixture
def tdu() -> TduTariff:
    return TduTariff(effective=dt.date(2023, 1, 1), fixed_usd_month=4.06, volumetric_ckwh=6.12)


@pytest.fixture
def plans() -> dict[str, Plan]:
    current = Plan(
        id="pulse_current",
        retailer="Pulse Power",
        name="Your Current Plan",
        term_months=12,
        base_charge_usd=4.95,
        energy_rates=[EnergyRate(rate_ckwh=15.8)],
        buyback=Buyback(kind=BuybackKind.fixed, rate_ckwh=15.8, offset_scope="energy_only"),
        etf_usd=20.0,
        etf_per_month_remaining=True,
        source="report-2026-07",
    )
    cheap = Plan(
        id="cheap_flat",
        retailer="Cheap Co",
        name="Flat Saver",
        term_months=12,
        base_charge_usd=0.0,
        energy_rates=[EnergyRate(rate_ckwh=8.0)],
        buyback=Buyback(kind=BuybackKind.none),
        source="manual",
    )
    needs_review = Plan(
        id="draft_plan",
        retailer="Draft Retailer",
        name="Uncertain Plan",
        term_months=24,
        base_charge_usd=9.95,
        energy_rates=[EnergyRate(rate_ckwh=20.0)],
        buyback=Buyback(kind=BuybackKind.none),
        needs_review=True,
        notes="auto-parsed, verify",
    )
    return {p.id: p for p in (current, cheap, needs_review)}


def test_build_workbook_smoke(intervals, tdu, plans):
    results = rank(list(plans.values()), intervals, tdu)
    assert len(results) == 3  # none of these plans need RTW prices

    buf = io.BytesIO()
    build_workbook(results, plans, intervals, tdu, buf)
    buf.seek(0)

    wb = openpyxl.load_workbook(buf, data_only=True)
    assert wb.sheetnames == ["Summary", "Monthly Detail", "Usage", "Plan Inputs"]

    # --- Summary sheet ---------------------------------------------------
    summary = wb["Summary"]
    header_row = 4  # 1-indexed: row 4 in openpyxl matches header_row=3 (0-idx)
    headers = [c.value for c in summary[header_row]]
    assert headers[:3] == ["Retailer", "Plan", "Term (mo)"]
    assert "1st-Year Net Bill" in headers
    assert "Import ¢/kWh (+TDU)" in headers

    net_bill_col = headers.index("1st-Year Net Bill") + 1
    plan_col_idx = headers.index("Plan") + 1

    by_plan_name_to_bill = {}
    r = header_row + 1
    while summary.cell(row=r, column=1).value not in (None, ""):
        name = summary.cell(row=r, column=plan_col_idx).value
        bill = summary.cell(row=r, column=net_bill_col).value
        by_plan_name_to_bill[name] = bill
        r += 1

    expected_by_plan_id = {res.plan_id: res.first_year_net for res in results}
    # Match by looking up each result's plan name (uses_rtw would add a
    # marker suffix, but none of these plans use RTW).
    for res in results:
        plan = plans[res.plan_id]
        assert by_plan_name_to_bill[plan.name] == pytest.approx(res.first_year_net, abs=1e-6)

    cheapest = min(results, key=lambda r: r.first_year_net)
    assert expected_by_plan_id[cheapest.plan_id] == min(expected_by_plan_id.values())

    # --- Monthly Detail sheet --------------------------------------------
    monthly = wb["Monthly Detail"]
    md_headers = [c.value for c in monthly[3]]
    assert md_headers == [
        "Plan ID",
        "Retailer",
        "Plan Name",
        "Month",
        "Import kWh",
        "Export kWh",
        "Energy Cost",
        "Base",
        "TDU",
        "Bill Credit",
        "Credit Earned",
        "Credit Used",
        "Rollover Out",
        "Bill",
    ]
    # 3 plans x 2 months = 6 data rows
    data_rows = 0
    r = 4
    while monthly.cell(row=r, column=1).value not in (None, ""):
        data_rows += 1
        r += 1
    assert data_rows == 6

    # Spot-check: sum of monthly "Bill" for pulse_current equals its first_year_net.
    pulse_result = next(res for res in results if res.plan_id == "pulse_current")
    bill_col = md_headers.index("Bill") + 1
    plan_id_col = md_headers.index("Plan ID") + 1
    total = 0.0
    r = 4
    while monthly.cell(row=r, column=1).value not in (None, ""):
        if monthly.cell(row=r, column=plan_id_col).value == "pulse_current":
            total += monthly.cell(row=r, column=bill_col).value
        r += 1
    assert total == pytest.approx(pulse_result.first_year_net, abs=1e-6)

    # --- Usage sheet -------------------------------------------------------
    usage = wb["Usage"]
    usage_headers = [c.value for c in usage[3]]
    assert usage_headers == ["Month", "Import kWh", "Export kWh", "Net kWh (import-export)"]
    # Two months of data
    assert usage.cell(row=4, column=1).value is not None
    assert usage.cell(row=5, column=1).value is not None

    # --- Plan Inputs sheet ---------------------------------------------------
    plan_inputs = wb["Plan Inputs"]
    pi_headers = [c.value for c in plan_inputs[3]]
    assert "id" in pi_headers
    assert "retailer" in pi_headers
    ids_col = pi_headers.index("id") + 1
    ids_seen = set()
    r = 4
    while plan_inputs.cell(row=r, column=1).value not in (None, ""):
        ids_seen.add(plan_inputs.cell(row=r, column=ids_col).value)
        r += 1
    assert ids_seen == {"pulse_current", "cheap_flat", "draft_plan"}


def test_build_workbook_to_path(tmp_path, intervals, tdu, plans):
    results = rank(list(plans.values()), intervals, tdu)
    out_path = tmp_path / "export.xlsx"
    build_workbook(results, plans, intervals, tdu, out_path)
    assert out_path.exists() and out_path.stat().st_size > 0

    wb = openpyxl.load_workbook(out_path)
    assert "Summary" in wb.sheetnames


def test_build_workbook_handles_empty_results(intervals, tdu, plans):
    """An empty results list (e.g. all plans skipped) should still produce a
    valid workbook rather than raising."""
    buf = io.BytesIO()
    build_workbook([], plans, intervals, tdu, buf)
    buf.seek(0)
    wb = openpyxl.load_workbook(buf)
    assert wb.sheetnames == ["Summary", "Monthly Detail", "Usage", "Plan Inputs"]
